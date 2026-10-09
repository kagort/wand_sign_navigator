#!/usr/bin/env python3
"""Собирает набор данных путеводителя из карточек content/ (frontmatter + ссылки).

Источник правды — карточки. Результат генерируется и руками не правится:
    data/people.json    персоналии: биография, портрет, разделы книги
    data/cards.json     все карточки (понятия, метафоры, кейсы, контекст, персоналии)
    data/sections.json  разделы глав: аннотация и карточки
    data/edges.json     связи: учитель/ученик/коллега, карточка→раздел, карточка→карточка
    data/people.xlsx    таблица персоналий в формате базы GVG (для просмотра)

Запуск:  python scripts/build_dataset.py
Нужны пакеты: pyyaml, openpyxl.
"""
import json
import os
import re
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / "content"
OUT = ROOT / "data"
SITE = "https://kagort.github.io/wand_sign_navigator/"

FOLDER_KIND = {
    "01-concepts": "концепт",
    "02-metaphors": "метафора",
    "03-people": "персоналия",
    "04-cases": "пример-кейс",
    "06-context": "контекст",
}
SECTION_RE = re.compile(r"\[\[(chapter_(\d)_([ivx]+)|predislovie)(?:[|\\\]#])")
CHAPTER_RE = re.compile(r"\[\[(?:chapter_(\d)(?:_[ivx]+)?|(predislovie))(?:[|\\\]#])")
LINK_RE = re.compile(r"(?<!!)\[\[([^\]|#\\]+)")
PORTRAIT_RE = re.compile(r"!\[\[(assets/[^\]|]+)")


def read_card(path):
    text = path.read_text(encoding="utf8")
    if not text.startswith("---\n"):
        return {}, text
    end = text.index("\n---\n", 4)
    return yaml.safe_load(text[4:end]) or {}, text[end + 5:]


def plain(md):
    """Markdown → короткий плоский текст для таблиц и всплывающих панелей."""
    md = re.sub(r"!\[\[[^\]]*\]\]", "", md)
    md = re.sub(r"\[\[[^\]|]*[|\\]+\|?([^\]]*)\]\]", r"\1", md)
    md = re.sub(r"\[\[([^\]]*)\]\]", r"\1", md)
    md = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", md)
    md = re.sub(r"^\s*>\s?", "", md, flags=re.M)
    md = re.sub(r"[*_`]+", "", md)
    return re.sub(r"\s+", " ", md).strip()


def section_text(body, heading):
    m = re.search(r"^##\s+" + re.escape(heading) + r".*?$\n(.*?)(?=^##\s|\Z)", body, re.M | re.S)
    return plain(m.group(1)) if m else ""


def link_id(value):
    m = re.match(r"\[\[([^\]|#]+)", str(value))
    return m.group(1).strip() if m else None


def year(value):
    m = re.search(r"(\d{3,4})", str(value or ""))
    if not m:
        return None
    y = int(m.group(1))
    return -y if "до н" in str(value) else y


def slug(path):
    return path.relative_to(CONTENT).with_suffix("").as_posix()


def main():
    cards, sections = {}, {}
    for path in sorted(CONTENT.rglob("*.md")):
        rel = path.relative_to(CONTENT)
        fm, body = read_card(path)
        cid = path.stem
        if cid == "predislovie":
            sections[cid] = dict(id=cid, chapter=0, section="", title="Предисловие", annotation=plain(re.sub(r"^#.*$", "", body, flags=re.M)), url=SITE + slug(path), cards=[])
            continue
        if rel.parts[0] == "05-chapters" and fm.get("type") == "раздел-главы":
            ch, sec = re.match(r"chapter_(\d)_([ivx]+)", cid).groups()
            sections[cid] = dict(
                id=cid, chapter=int(ch), section=sec,
                title=re.sub(r"^Гл\. \d \([ivx]+\)\s*", "", str(fm.get("title", ""))),
                annotation=section_text(body, "Аннотация"),
                url=SITE + slug(path), cards=[],
                _listed=set(LINK_RE.findall(body.split("## Карточки раздела", 1)[-1])),
            )
            continue
        kind = FOLDER_KIND.get(rel.parts[0])
        if not kind or fm.get("draft") is True:
            continue
        portrait = PORTRAIT_RE.search(body)
        cards[cid] = dict(
            id=cid, kind=kind, folder=rel.parts[0],
            title=str(fm.get("title", cid)), title_en=fm.get("title_en"),
            status=fm.get("status"), tags=fm.get("tags") or [],
            url=SITE + slug(path),
            portrait=(SITE + rel.parts[0] + "/" + portrait.group(1)) if portrait else None,
            sections=sorted({m.group(1) for m in SECTION_RE.finditer(body)}),
            chapters=sorted({int(m.group(1) or 0) for m in CHAPTER_RE.finditer(body)}),
            links=sorted({l.strip() for l in LINK_RE.findall(body + json.dumps(fm.get("related") or [], ensure_ascii=False))}),
            _fm=fm, _body=body,
        )

    for c in cards.values():
        c["links"] = [l for l in c["links"] if l in cards and l != c["id"]]
    # карточка относится к разделу, если ссылается на него ИЛИ перечислена на странице раздела
    for sid, s in sections.items():
        for cid in s.pop("_listed", set()):
            if cid in cards and sid not in cards[cid]["sections"]:
                cards[cid]["sections"].append(sid)
    for c in cards.values():
        c["sections"].sort()
        for s in c["sections"]:
            if s in sections:
                sections[s]["cards"].append(c["id"])

    people, edges = [], []
    for c in cards.values():
        if c["kind"] != "персоналия":
            continue
        fm, body = c["_fm"], c["_body"]
        rel = {k: [i for i in (link_id(v) for v in fm.get(k) or []) if i] for k in ("teachers", "students", "colleagues")}
        for t in rel["teachers"]:
            edges.append(dict(source=t, target=c["id"], type="учитель→ученик"))
        for t in rel["students"]:
            edges.append(dict(source=c["id"], target=t, type="учитель→ученик"))
        for t in rel["colleagues"]:
            edges.append(dict(source=c["id"], target=t, type="коллеги"))
        people.append(dict(
            id=c["id"], title=c["title"], title_en=c["title_en"], url=c["url"], portrait=c["portrait"],
            years=fm.get("years"), born=fm.get("born"), died=fm.get("died"),
            born_year=year(fm.get("born") or fm.get("years")), died_year=year(fm.get("died")),
            birth_place=fm.get("birth_place"), birth_coords=fm.get("birth_coords"),
            death_place=fm.get("death_place"), death_coords=fm.get("death_coords"),
            nationality=fm.get("nationality"), alma_mater=fm.get("alma_mater"),
            affiliations=fm.get("affiliations") or [], fields=fm.get("fields") or [],
            awards=fm.get("awards") or [], sep=fm.get("sep"), wiki_ru=fm.get("wiki_ru"),
            notes=fm.get("notes"), bio_checked=bool(fm.get("bio_checked")),
            tags=c["tags"], sections=c["sections"], chapters=c["chapters"], **rel,
            who=section_text(body, "Кто это"),
            role=section_text(body, "Роль в аргументации Уилсона"),
            works=section_text(body, "Ключевые понятия/работы"),
        ))

    # дубли рёбер (A коллега B и B коллега A) схлопываем
    seen, uniq = set(), []
    for e in edges:
        key = (e["type"],) + (tuple(sorted((e["source"], e["target"]))) if e["type"] == "коллеги" else (e["source"], e["target"]))
        if key not in seen and e["source"] in cards and e["target"] in cards:
            seen.add(key)
            uniq.append(e)
    for c in cards.values():
        uniq += [dict(source=c["id"], target=s, type="упомянут-в-разделе") for s in c["sections"] if s in sections]
        uniq += [dict(source=c["id"], target=l, type="ссылка") for l in c["links"]]

    OUT.mkdir(exist_ok=True)
    clean = [{k: v for k, v in c.items() if not k.startswith("_")} for c in cards.values()]
    order = lambda s: (s["chapter"], ["", "i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi"].index(s["section"]))
    dump = lambda name, data: (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf8")
    dump("cards.json", clean)
    dump("people.json", people)
    dump("sections.json", sorted(sections.values(), key=order))
    dump("edges.json", uniq)
    write_xlsx(people, OUT / "people.xlsx")
    print(f"карточек {len(clean)}, персоналий {len(people)}, разделов {len(sections)}, связей {len(uniq)} → {OUT.relative_to(ROOT)}/")


GROUPS = [  # (название группы, цвет, [(колонка, функция)])
    ("ИДЕНТИФИКАЦИЯ", "FFE699", [
        ("Имя и фамилия", lambda p: p["title"]), ("Имя (оригинал)", lambda p: p["title_en"]),
        ("Национальность", lambda p: p["nationality"]), ("Портрет", lambda p: p["portrait"])]),
    ("БИОГРАФИЯ — ВРЕМЯ", "C6E0B4", [
        ("Годы жизни", lambda p: p["years"]), ("Дата рождения", lambda p: p["born"]), ("Дата смерти", lambda p: p["died"])]),
    ("БИОГРАФИЯ — МЕСТО", "D9C3E9", [
        ("Место рождения", lambda p: p["birth_place"]), ("Координаты рождения", lambda p: coords(p["birth_coords"])),
        ("Место смерти", lambda p: p["death_place"]), ("Координаты смерти", lambda p: coords(p["death_coords"]))]),
    ("АКАДЕМИЧЕСКАЯ КАРЬЕРА", "F4B6B6", [
        ("Alma mater", lambda p: p["alma_mater"]), ("Аффилиация", lambda p: "; ".join(p["affiliations"])),
        ("Учителя", lambda p: names(p["teachers"])), ("Ученики", lambda p: names(p["students"])),
        ("Коллеги", lambda p: names(p["colleagues"])), ("Награды и премии", lambda p: "; ".join(p["awards"]))]),
    ("НАУЧНЫЙ ПРОФИЛЬ", "B4E5E4", [
        ("Области интересов", lambda p: "; ".join(p["fields"])), ("Кто это", lambda p: p["who"]),
        ("Основные труды / понятия", lambda p: p["works"]), ("Ссылка SEP", lambda p: p["sep"]),
        ("Ссылка Wikipedia (RU)", lambda p: p["wiki_ru"])]),
    ("КОНТЕКСТ КНИГИ УИЛСОНА", "BDD7EE", [
        ("Роль в книге", lambda p: p["role"]), ("Главы", lambda p: "; ".join("предисловие" if c == 0 else f"гл. {c}" for c in p["chapters"])),
        ("Разделы книги", lambda p: "; ".join(p["sections"])),
        ("Карточка", lambda p: p["url"]), ("Данные проверены", lambda p: "да" if p["bio_checked"] else "нет"),
        ("Примечания", lambda p: p["notes"])]),
]
TITLES = {}


def coords(c):
    return f"{c[0]:.4f}, {c[1]:.4f}" if c else None


def names(ids):
    return "; ".join(TITLES.get(i, i) for i in ids)


def write_xlsx(people, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    TITLES.update({p["id"]: p["title"] for p in people})
    wb = Workbook()
    ws = wb.active
    ws.title = "База данных"
    ws["A1"] = "БАЗА ДАННЫХ ПЕРСОНАЛИЙ · «Блуждающие значения» М. Уилсона (генерируется из карточек — не править)"
    ws["A1"].font = Font(bold=True, size=13)
    col = 1
    for group, color, cols in GROUPS:
        fill = PatternFill("solid", fgColor=color)
        ws.cell(row=2, column=col, value=group).font = Font(bold=True)
        for name, _ in cols:
            for r in (2, 3):
                ws.cell(row=r, column=col).fill = fill
            ws.cell(row=3, column=col, value=name).font = Font(bold=True)
            col += 1
    flat = [f for _, _, cols in GROUPS for f in cols]
    for r, p in enumerate(sorted(people, key=lambda p: (p["born_year"] is None, p["born_year"] or 0)), start=4):
        for c, (_, fn) in enumerate(flat, start=1):
            v = fn(p)
            ws.cell(row=r, column=c, value=v if v not in ("", []) else None).alignment = Alignment(wrap_text=True, vertical="top")
    for c, (name, _) in enumerate(flat, start=1):
        ws.column_dimensions[get_column_letter(c)].width = 60 if name in ("Кто это", "Роль в книге", "Основные труды / понятия") else 22
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(flat))}{3 + len(people)}"

    ref = wb.create_sheet("Справочник полей")
    ref.append(["Колонка", "Поле во frontmatter карточки", "Формат"])
    for row in [
        ("Дата рождения / смерти", "born / died", "ГГГГ-ММ-ДД, ГГГГ или «ок. ГГГГ»"),
        ("Место рождения / смерти", "birth_place / death_place", "текст"),
        ("Координаты", "birth_coords / death_coords", "[широта, долгота]"),
        ("Национальность", "nationality", "текст"),
        ("Alma mater", "alma_mater", "текст"),
        ("Аффилиация, области интересов, награды", "affiliations, fields, awards", "список"),
        ("Учителя / ученики / коллеги", "teachers / students / colleagues", "список ссылок [[id]] на карточки"),
        ("Ссылки SEP / Wikipedia", "sep / wiki_ru", "URL"),
        ("Данные проверены", "bio_checked", "true / false"),
        ("Кто это, роль в книге, труды", "разделы карточки «Кто это», «Роль в аргументации Уилсона», «Ключевые понятия/работы»", "текст карточки"),
        ("Разделы книги", "ссылки [[chapter_N_x]] в карточке", "вычисляется"),
    ]:
        ref.append(row)
    for c in "ABC":
        ref.column_dimensions[c].width = 45
    wb.save(path)


if __name__ == "__main__":
    os.chdir(ROOT)
    main()
